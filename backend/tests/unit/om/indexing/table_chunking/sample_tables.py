"""Sample tables for table-chunking tests.

The two department CSVs are copies of the synthetic (mock) gov test data in
``tests/workflow_creator/test_data`` (aarogyasri AR1 preauth, cid CID1
depositors and bank statements). They are kept here as strings because
``*.csv`` files are git-ignored in this repo.
"""

import csv
import io
from collections.abc import Callable
from collections.abc import Iterator
from functools import partial

from om.indexing.table_chunking.sheet_summary import SheetRows

AAROGYASRI_PREAUTH_CSV = """field,value
preauth_id,PA-AR-2026-1001
beneficiary_card,ASRI-XXXX-1001 (mock)
patient,Patient A (mock; 58M)
diagnosis,Acute ST-elevation myocardial infarction (STEMI)
icd_code,I21.0
proposed_package_code,AS-CARD-012
proposed_package_name,PTCA with single drug-eluting stent
specialty,Cardiology
cost_estimate,90000
investigations_uploaded,ECG; Troponin-I; Coronary Angiogram
clinical_justification,Single-vessel LAD occlusion on angiogram; primary PCI indicated
emergency,NO
doc_completeness,COMPLETE
"""

CID_DEPOSITORS_CSV = """depositor_id,name,amount_inr,deposit_date,receipt_no,branch
D-1001,Ramesh Kumar,500000,2025-11-02,SF-1001,Ameerpet
D-1002,Ramesh Kumr,500000,2025-11-02,SF-1044,Ameerpet
D-1003,Lakshmi Devi,300000,2025-11-10,SF-1102,Dilsukhnagar
D-1004,Syed Imran,750000,2025-12-01,SF-1210,Ameerpet
D-1005,Venkata Rao,250000,2025-12-05,SF-1288,Kukatpally
"""

CID_BANK_STATEMENTS_CSV = """txn_id,from_account,to_account,amount_inr,date,note
T-01,Depositors (various),SF-COLLECT-50100111 HDFC,1800000,2025-12-06,Deposits pooled into collection account
T-02,SF-COLLECT-50100111 HDFC,BRIGHT-VENTURES-50200222 ICICI,1500000,2025-12-08,Transfer to shell entity (layer 1)
T-03,BRIGHT-VENTURES-50200222 ICICI,APEX-HOLDINGS-50300333 AXIS,1200000,2025-12-11,Onward layering to second shell
T-04,APEX-HOLDINGS-50300333 AXIS,Property purchase (villa Kompally),900000,2025-12-20,Exit into real estate (attachable)
"""

DISTRICTS = ["Hyderabad", "Warangal", "Karimnagar", "Nizamabad", "Khammam"]
STATUSES = ["Pending", "Resolved", "Rejected"]


def grievance_rows(count: int, start: int = 0) -> list[list[str]]:
    """A large register: id, district, status, amount, date, free-text remark."""
    rows = [["case_id", "district", "status", "amount", "filed_on", "remark"]]
    for i in range(start, start + count):
        rows.append(
            [
                f"GRV-{i:05d}",
                DISTRICTS[i % len(DISTRICTS)],
                STATUSES[i % len(STATUSES)],
                str(100 + (i % 50) * 10),
                f"2025-{1 + i % 6:02d}-{1 + i % 28:02d}",
                f"Remark number {i} about drainage near ward {i % 97}",
            ]
        )
    return rows


def rows_to_csv(rows: list[list[str]]) -> str:
    out = io.StringIO()
    csv.writer(out, lineterminator="\n").writerows(rows)
    return out.getvalue()


def _iter_rows(rows: list[list[str]]) -> Iterator[list[str]]:
    return iter([list(row) for row in rows])


def sheet(name: str, rows: list[list[str]]) -> SheetRows:
    opener: Callable[[], Iterator[list[str]]] = partial(_iter_rows, rows)
    return SheetRows(name=name, open_rows=opener)


def csv_sheet(name: str, csv_text: str) -> SheetRows:
    rows = list(csv.reader(io.StringIO(csv_text)))
    return sheet(name, rows)


def build_sample_xlsx() -> bytes:
    """A 3-sheet workbook built in memory (nothing binary is committed):
    Readme (text), Register (data with a quoted cell, a newline and a 0) and an
    empty sheet."""
    import openpyxl

    workbook = openpyxl.Workbook()
    readme = workbook.active
    assert readme is not None
    readme.title = "Readme"
    readme.append(["Grievance register, mock data"])
    readme.append(["Maintained by the district office."])
    register = workbook.create_sheet("Register")
    register.append(["case_id", "district", "amount", "note"])
    register.append(["GRV-1", "Hyderabad", 0, 'said "urgent", twice'])
    register.append([None, None, None, None])
    register.append(["GRV-2", "Warangal", 4500, "line one\nline two"])
    workbook.create_sheet("Empty")
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
